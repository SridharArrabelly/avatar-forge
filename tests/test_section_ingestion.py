"""Offline section-ingestion, immutable-index and upload-result contracts."""
from __future__ import annotations

import os
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import setup_aisearch_index as setup
from backend.document_sections import catalogue_title


ENV = {
    "AZURE_SEARCH_ENDPOINT": "https://example.invalid",
    "SEARCH_INDEX_NAME": "minutes-v2",
    "PROJECT_ENDPOINT": "https://example.invalid/api/projects/test",
}


def rejects(call, expected):
    try:
        call()
    except (ValueError, RuntimeError) as error:
        assert expected in str(error), str(error)
    else:
        raise AssertionError(f"Expected rejection containing {expected!r}")


def main() -> int:
    assert catalogue_title("Board Meeting 15 March 2006 | Header") == "Board Meeting 15 March 2006"
    assert catalogue_title("Board Meeting 15 March 2006") == "Board Meeting 15 March 2006"
    assert catalogue_title("Strategy | Risk Review") == "Strategy | Risk Review"
    with patch.dict(os.environ, ENV, clear=True):
        defaults = setup.load_settings()
        assert defaults["chunking_mode"] == "section" and defaults["document_scope"] == "minutes"
    with patch.dict(os.environ, {**ENV, "CHUNKING_MODE": "window", "DOCUMENT_SCOPE": "all"}, clear=True):
        legacy = setup.load_settings()
        assert legacy["chunking_mode"] == "window" and legacy["document_scope"] == "all"
    for values, message in (
        ({"CHUNKING_MODE": "invalid"}, "CHUNKING_MODE"),
        ({"DOCUMENT_SCOPE": "invalid"}, "DOCUMENT_SCOPE"),
        ({"CHUNKING_MODE": "section", "DOCUMENT_SCOPE": "all"}, "DOCUMENT_SCOPE=minutes"),
        ({"CHUNKING_MODE": "section", "DOCUMENT_SCOPE": "minutes", "RECREATE_INDEX": "true"}, "versioned"),
    ):
        with patch.dict(os.environ, {**ENV, **values}, clear=True):
            rejects(setup.load_settings, message)
    print("PASS section/minutes defaults, explicit window/all compatibility, and invalid settings")

    section = {**defaults, "chunking_mode": "section", "document_scope": "minutes", "embed_dim": 1536}
    schema = setup.build_index("minutes-v2", section)
    assert setup.SECTION_FIELDS <= {field.name for field in schema.fields}
    assert schema.semantic_search.default_configuration_name == section["semantic_config"]
    old_schema = setup.build_index("minutes-v2", legacy)
    assert not setup.SECTION_FIELDS & {field.name for field in old_schema.fields}
    assert old_schema.semantic_search.default_configuration_name is None
    for settings, existing in ((section, old_schema), (legacy, schema)):
        client = MagicMock()
        client.list_indexes.return_value = [existing]
        with patch.object(setup, "make_index_client", return_value=client):
            rejects(lambda: setup.ensure_index(dict(settings)), "cannot be changed in place")
        client.delete_index.assert_not_called()
        client.create_or_update_index.assert_not_called()
    client = MagicMock()
    client.list_indexes.return_value = [schema]
    with patch.object(setup, "make_index_client", return_value=client):
        retained = dict(section)
        setup.ensure_index(retained)
        assert retained["section_index_exists"]
    client.create_or_update_index.assert_not_called()
    print("PASS section schema is explicit and existing layout/index versions are preserved")

    with TemporaryDirectory() as directory:
        root = Path(directory)
        source = root / "Board Meeting 15 March 2006.docx"
        doc = Document()
        for text in [
            source.stem, "Attendees: Alice; Bob.",
            "Agenda:", "Review the synthetic test.",
            "Key Discussion Points:", "Discussed source integrity.",
            "Decisions Made:", "Approved a limited test.",
            "Action Items:", "Alice: check the evidence.",
            "Next Steps:", "Review the results.",
        ]:
            doc.add_paragraph(text)
        doc.save(source)
        policy = root / "policies"
        policy.mkdir()
        (policy / "do-not-read.docx").write_bytes(b"not a document; must be excluded")
        settings = {**section, "data_dir": root}
        documents = setup.prepare_section_documents(settings)
        assert len(documents) == 6
        assert documents[0]["section"] == "Header" and documents[0]["chunk_index"] == 0
        assert all(document["documentType"] == "MeetingMinutes" for document in documents)
        assert len({document["id"] for document in documents}) == 6
        assert all(document["source_sha256"] for document in documents)
        doc.paragraphs[-1]._p.append(OxmlElement("w:footnoteReference"))
        doc.save(source)
        rejects(lambda: setup.prepare_section_documents(settings), "footnoteReference")
    print("PASS whole sections preserve context and never open excluded policy documents")

    ready = {**section, "section_documents": documents}
    search = MagicMock()
    search.__enter__.return_value = search
    search.search.side_effect = [[], documents]
    with patch.object(setup, "make_search_client", return_value=search), patch.object(setup.time, "sleep") as sleep:
        setup.verify_section_documents(ready)
        sleep.assert_called_once_with(2)
    search.search.side_effect = None
    search.search.return_value = [{**documents[0], "content": "Altered"}]
    with patch.object(setup, "make_search_client", return_value=search):
        rejects(lambda: setup.verify_section_documents(ready, timeout_s=0), "differs")
    print("PASS exact readback tolerates indexing visibility delay, not changed content")

    search = MagicMock()
    small = [{"id": "a"}, {"id": "b"}]
    search.upload_documents.return_value = [
        SimpleNamespace(key="a", succeeded=True, error_message=None),
        SimpleNamespace(key="b", succeeded=False, error_message="rejected"),
    ]
    with patch.object(setup, "make_search_client", return_value=search):
        rejects(lambda: setup.upload({}, iter(small)), "upload failed")
    client = MagicMock()
    client.embeddings.create.return_value = SimpleNamespace(data=[
        SimpleNamespace(index=1, embedding=[2.0]),
        SimpleNamespace(index=0, embedding=[1.0]),
    ])
    assert setup.embed_batch(client, "embedding", ["first", "second"]) == [[1.0], [2.0]]
    print("PASS per-document upload failures and embedding order cannot silently corrupt the index")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
